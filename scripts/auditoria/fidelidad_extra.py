"""Fidelidad de las capas de la sesión 4 frente a SUS fuentes (no frente al capítulo):

  - Versión extendida: cada parte, literal en su borrador (V93, V85, V79 o tablas V85), con las
    convenciones de la app (DM1, UI, «duración de la insulina activa»).
  - Para el paciente: la información para pacientes, idéntica a la V5 del autor; el resumen,
    idéntico a la maquetación de la editorial (y a la V6 del autor).
  - Test: cada cita del capítulo, en su página del PDF final.

Uso (desde la raíz, tras `npx vite-node scripts/auditoria/exportar-contenido.mjs`):
    python scripts/auditoria/fidelidad_extra.py
Las rutas de los documentos del autor están abajo; se pueden pasar otras con variables de
entorno (MSEEN_V93, MSEEN_V85, MSEEN_V79, MSEEN_T85, MSEEN_V5, MSEEN_RESUMEN_PDF, MSEEN_PDF).
Resultado en _audit_fidelidad_extra.json (ignorado por git).
"""
import json
import os
import re
import sys
import unicodedata

import docx
import fitz

H = r"C:\Users\giorg\OneDrive"
RUTAS = {
    "V93": os.environ.get("MSEEN_V93", H + r"\Capitulo SEEN\Definitivo\Capitulo_SEEN_AID_V93_integrado.docx"),
    "V85": os.environ.get("MSEEN_V85", H + r"\Capitulo SEEN\30 mayo\Capitulo_SEEN_AID_V85_limpio.docx"),
    "V79": os.environ.get("MSEEN_V79", H + r"\Capitulo SEEN\Capitulo_SEEN_AID_V79_limpio.docx"),
    "T85": os.environ.get("MSEEN_T85", H + r"\Capitulo SEEN\30 mayo\Tablas_SEEN_AID_V85_vertical.docx"),
    "V5": os.environ.get("MSEEN_V5", H + r"\Desktop\Capitulo SEEN\Informacion_pacientes_AID_SEEN_version_definitiva_V5.docx"),
    "PDF": os.environ.get("MSEEN_PDF", r"C:\Users\giorg\Downloads\Capitulo_AID_SEEN_05-10-2026_ANOTADO_3_DETALLES_FINALES.pdf"),
}


def resumen_pdf():
    """La maquetación del resumen está en el Escritorio con el nombre en NFD."""
    if os.environ.get("MSEEN_RESUMEN_PDF"):
        return os.environ["MSEEN_RESUMEN_PDF"]
    d = H + r"\Desktop"
    for f in os.listdir(d):
        if unicodedata.normalize("NFC", f).endswith("_Resumen.pdf") and "Tratamiento" in f:
            return os.path.join(d, f)
    return None


def plano(s):
    s = unicodedata.normalize("NFC", s)
    s = s.replace("\u00ad", "").replace("“", "«").replace("”", "»").replace("—", "-").replace("–", "-")
    return re.sub(r"\s+", " ", s).strip()


def conv(s):
    s = s.replace("DT1", "DM1").replace("DT2", "DM2")
    s = re.sub(r"(\d) U/día", r"\1 UI/día", s)
    s = s.replace("una AIT de 2 h", "una duración de la insulina activa de 2 h")
    s = s.replace("AIT 2 h", "duración de la insulina activa 2 h")
    s = s.replace("duración de insulina activa", "duración de la insulina activa")
    return s


def texto_docx(ruta):
    d = docx.Document(ruta)
    trozos = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                trozos.append(c.text)
    return plano(conv(" \n ".join(trozos)))


def paginas_pdf(ruta):
    doc = fitz.open(ruta)
    out = []
    for p in doc:
        t = p.get_text()
        t = re.sub(r"-\n(?=[a-záéíóúñ])", "", t)
        out.append(plano(t))
    return out


app = json.load(open("_audit_contenido.json", encoding="utf-8"))
res = {"extendida": [], "pacientes": [], "resumen": [], "test": []}

# 1. Versión extendida
fuentes = {k: texto_docx(RUTAS[k]) for k in ("V93", "V85", "V79", "T85") if os.path.exists(RUTAS[k])}
ok_ext = 0
for f in app.get("extendida", []):
    for p in f["partes"]:
        src = fuentes.get(p["borrador"])
        if src is None:
            res["extendida"].append({"id": f["id"], "motivo": f"no encuentro {p['borrador']}"})
        elif plano(p["texto"]) not in src:
            res["extendida"].append({"id": f["id"], "borrador": p["borrador"], "texto": p["texto"][:120]})
        else:
            ok_ext += 1

# 2. Información para pacientes = V5
v5 = [plano(x) for x in texto_docx(RUTAS["V5"]).split(" \n ") if x.strip()] if os.path.exists(RUTAS["V5"]) else []
v5_todo = " ".join(v5)
ok_pac = 0
for s in app["pacientes"]["informacion"]:
    for t in [s["pregunta"], *s["parrafos"]]:
        if plano(t) in v5_todo:
            ok_pac += 1
        else:
            res["pacientes"].append(t[:120])

# 3. Resumen = maquetación (comparación por palabras, sin cortes de línea ni guiones)
rp = resumen_pdf()
ok_res = 0
if rp:
    pdf_res = plano(" ".join(paginas_pdf(rp)))
    pdf_res = re.sub(r"ec-europe - \d\d/\d\d/\d{4} \d+ ", "", pdf_res)
    for t in app["pacientes"]["resumen"]:
        if plano(t) in pdf_res:
            ok_res += 1
        else:
            res["resumen"].append(t[:120])
else:
    res["resumen"].append("no encuentro el PDF del resumen")

# 4. Test: citas en su página del capítulo
pags = paginas_pdf(RUTAS["PDF"]) if os.path.exists(RUTAS["PDF"]) else []
ok_test = 0
for q in app.get("test", []):
    for c in q["citas"]:
        for trozo in c["texto"].split(" […] "):
            t = plano(trozo)
            p = c["p"]
            if any(t in pags[i] for i in range(max(0, p - 2), min(len(pags), p + 1))):
                ok_test += 1
            else:
                res["test"].append({"id": q["id"], "p": p, "texto": trozo[:120]})

json.dump(res, open("_audit_fidelidad_extra.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
sys.stdout.reconfigure(encoding="utf-8")
print(f"versión extendida: {ok_ext} partes literales en su borrador; fallos {len(res['extendida'])}")
print(f"información para pacientes: {ok_pac} frases idénticas a la V5; fallos {len(res['pacientes'])}")
print(f"resumen: {ok_res} párrafos idénticos a la maquetación; fallos {len(res['resumen'])}")
print(f"test: {ok_test} citas en su página del PDF; fallos {len(res['test'])}")
for k, v in res.items():
    for x in v[:5]:
        print("  ", k, x)
